#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""tests/oracle/xlsx/generar.py — (2026.12.31) los libros de Excel de prueba del lector propio (core/data/xlsx/lector.js),
escritos por otros programas: openpyxl, XlsxWriter y LibreOffice (por UNO), más unos escritos a mano para lo que esos
programas no escriben (prefijos de espacio de nombres, celdas sin referencia, guía fonética, escapes _xHHHH_, la variante
«strict», ZIP64 con descriptores de datos) y los formatos que el lector debe rechazar con un mensaje claro (.xls, .ods,
.xlsb, un libro cifrado y un documento de Word). Lo que contienen lo lee openpyxl en ../xlsx_oraculo.py.

Uso:  python3 generar.py          (hace falta openpyxl, XlsxWriter y LibreOffice con su módulo uno de Python)
No forma parte de la batería: los archivos ya están en el repositorio y solo hay que regenerarlos si cambia este script.
"""
import datetime as dt
import io
import os
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import time
import zipfile
import zlib

import openpyxl
import xlsxwriter
from openpyxl.utils.datetime import CALENDAR_MAC_1904

AQUI = os.path.dirname(os.path.abspath(__file__))
CREADO = dt.datetime(2026, 10, 20, 9, 30, 0)


def ruta(nombre):
    return os.path.join(AQUI, nombre)


# ---------------------------------------------------------------- openpyxl
def openpyxl_tipos():
    wb = openpyxl.Workbook()
    wb.properties.creator = 'StatSim Pro (pruebas)'
    wb.properties.created = CREADO
    ws = wb.active
    ws.title = 'Tipos'
    columnas = [
        ('Entero', None, [1, -2, 0, 123456789012345, None, 7]),
        ('Decimal', None, [0.1, -2.5, 1e-300, 9.99e307, 3.141592653589793, -0.0]),
        ('Texto', None, ['Ñandú ✓ 🎉', '  espacios  ', 'línea1\nlínea2', 'tab\tdentro', 'menor < y & mayor >', 'comillas "dobles" y \'simples\'']),
        ('Booleano', None, [True, False, True, None, False, True]),
        ('Fecha', 'dd/mm/yyyy', [dt.date(2024, 5, 1), dt.date(1900, 1, 1), dt.date(1900, 2, 28), dt.date(1900, 3, 1), dt.date(9999, 12, 31), dt.date(1970, 1, 1)]),
        ('FechaHora', 'yyyy-mm-dd hh:mm:ss', [dt.datetime(2024, 5, 1, 13, 30, 15), dt.datetime(2000, 2, 29, 23, 59, 59), dt.datetime(1999, 12, 31, 0, 0, 1), None, dt.datetime(2024, 1, 1), dt.datetime(1950, 6, 15, 6, 7, 8)]),
        ('Hora', 'hh:mm', [dt.time(13, 30), dt.time(0, 0, 1), dt.time(23, 59, 59), dt.time(12, 0), None, dt.time(6, 45, 30)]),
        ('Formula', None, ['=1+1', '=A2*2', '=SUM(A2:A3)', None, '=B2&"x"', '=NOW()']),
        ('Error', None, ['#N/A', '#DIV/0!', '#VALUE!', '#REF!', '#NAME?', '#NUM!']),
        ('Duracion', '[h]:mm:ss', [dt.timedelta(hours=36), dt.timedelta(minutes=90), dt.timedelta(seconds=1), None, dt.timedelta(days=2, hours=1), dt.timedelta(0)]),
        ('Porcentaje', '0.00%', [0.125, 1, -0.5, None, 0, 2.5]),
        ('TextoDia', '"día "0', [5, 6, 7, None, 8, 9]),
        ('Rojo', '[Red]0.00', [-3.5, 2, None, 4, 5, 6]),
        ('Cientifica', '0.00E+00', [12345.678, 1e-10, None, 2, 3, 4]),
        ('Euros', '#,##0.00\\ "€"', [1234.5, -0.01, None, 1e6, 2, 3]),
        ('MesAnio', 'mmm-yy', [dt.date(2024, 5, 1), dt.date(2001, 12, 1), None, dt.date(1990, 1, 1), dt.date(2030, 7, 1), dt.date(2024, 2, 29)]),
        ('General', 'General', [42, 'cuarenta y dos', None, 0, -1, 1.5]),
        ('Literal', '0" mm"', [1.5, 2, None, 3, 4, 5]),
        ('Escapada', '0\\ \\m', [10, 20, None, 30, 40, 50]),
    ]
    for j, (titulo, formato, valores) in enumerate(columnas, start=1):
        ws.cell(row=1, column=j, value=titulo)
        for i, v in enumerate(valores, start=2):
            if v is None:
                continue
            c = ws.cell(row=i, column=j, value=v)
            if formato:
                c.number_format = formato
    oculta = wb.create_sheet('Oculta')
    oculta['A1'] = 'solo'
    oculta['B2'] = 5
    oculta.sheet_state = 'hidden'
    muy = wb.create_sheet('MuyOculta')
    muy['A1'] = 1
    muy.sheet_state = 'veryHidden'
    disp = wb.create_sheet('Dispersa')
    disp['A1'] = 'a'
    disp['C3'] = 3
    disp['Z100'] = 'z'
    disp['AA1'] = 27
    disp['B50'].number_format = '0.00'   # una celda con formato y sin valor: no es un dato
    wb.create_sheet('Vacia')
    tit = wb.create_sheet('Titulo')
    tit['A1'] = 'Encuesta de satisfacción 2024'
    for j, h in enumerate(['ID', 'Edad', 'Sexo', 'Satisfaccion'], start=1):
        tit.cell(row=3, column=j, value=h)
    filas = [(1, 20, 'Mujer', 4), (2, 31, 'Hombre', 5), (3, None, 'Mujer', 3), (4, 45, None, 2), (5, 52, 'Hombre', None)]
    for i, f in enumerate(filas, start=4):
        for j, v in enumerate(f, start=1):
            if v is not None:
                tit.cell(row=i, column=j, value=v)
    wb.save(ruta('openpyxl_tipos.xlsx'))


def openpyxl_iso():
    wb = openpyxl.Workbook(iso_dates=True)
    wb.properties.created = CREADO
    ws = wb.active
    ws.title = 'Iso'
    ws.append(['Fecha', 'FechaHora', 'Numero'])
    ws.append([dt.date(2024, 5, 1), dt.datetime(2024, 5, 1, 13, 30, 15), 1])
    ws.append([dt.date(1999, 12, 31), dt.datetime(1999, 12, 31, 23, 59, 59), 2])
    ws.append([dt.date(2000, 2, 29), dt.datetime(2000, 1, 1, 0, 0, 0), 3])
    wb.save(ruta('openpyxl_iso.xlsx'))


def openpyxl_1904():
    wb = openpyxl.Workbook()
    wb.epoch = CALENDAR_MAC_1904
    wb.properties.created = CREADO
    ws = wb.active
    ws.title = 'Mac1904'
    ws.append(['Fecha', 'FechaHora', 'Numero'])
    for f, fh, n in [(dt.date(2024, 5, 1), dt.datetime(2024, 5, 1, 13, 30, 15), 1), (dt.date(1904, 1, 1), dt.datetime(1904, 1, 2, 6, 0, 0), 2), (dt.date(1985, 3, 4), dt.datetime(2010, 10, 10, 10, 10, 10), 3)]:
        ws.append([f, fh, n])
    for fila in ws.iter_rows(min_row=2):
        fila[0].number_format = 'yyyy-mm-dd'
        fila[1].number_format = 'yyyy-mm-dd hh:mm:ss'
    wb.save(ruta('openpyxl_1904.xlsx'))


# ---------------------------------------------------------------- XlsxWriter
def xlsxwriter_libro(nombre, opciones):
    wb = xlsxwriter.Workbook(ruta(nombre), opciones)
    wb.set_properties({'created': CREADO, 'author': 'StatSim Pro (pruebas)'})
    negrita = wb.add_format({'bold': True})
    fecha = wb.add_format({'num_format': 'dd/mm/yyyy'})
    fecha_hora = wb.add_format({'num_format': 'yyyy-mm-dd hh:mm'})
    duracion = wb.add_format({'num_format': '[h]:mm'})
    porcentaje = wb.add_format({'num_format': '0.0%'})
    rojo = wb.add_format({'font_color': 'red'})
    ws = wb.add_worksheet('Datos')
    encabezado = ['ID', 'Nombre', 'Grupo', 'Nota', 'Nacimiento', 'Cita', 'Tiempo', 'Aprobado', 'Calculo', 'Pendiente', 'Texto calc', 'Error calc', 'Logico calc', 'Avance', 'Rico', 'Enlace']
    for j, h in enumerate(encabezado):
        ws.write_string(0, j, h, negrita)
    nombres = ['Ana', 'Bruno', 'Carmen', 'Diego', 'Elena', 'Fermín', 'Gloria', 'Héctor']
    for i in range(8):
        f = i + 1
        ws.write_number(f, 0, f)
        ws.write_string(f, 1, nombres[i])
        ws.write_string(f, 2, ['A', 'B'][i % 2])                 # repetidas: una sola cadena compartida
        if i != 3:
            ws.write_number(f, 3, 10 + i * 1.25)
        ws.write_datetime(f, 4, dt.datetime(1990 + i, 1 + i, 1 + i), fecha)
        ws.write_datetime(f, 5, dt.datetime(2024, 3, 1 + i, 8 + i, 15), fecha_hora)
        ws.write_number(f, 6, (i + 1) / 24 * 7, duracion)        # horas acumuladas: una duración
        ws.write_boolean(f, 7, i % 3 != 0)
        ws.write_formula(f, 8, f'=A{f + 1}*2', None, f * 2)      # con su resultado guardado
        ws.write_formula(f, 9, f'=A{f + 1}+D{f + 1}')            # sin resultado: XlsxWriter guarda 0 y pide recalcular
        ws.write_formula(f, 10, f'=B{f + 1}&"!"', None, nombres[i] + '!')
        ws.write_formula(f, 11, '=1/0', None, '#DIV/0!')
        ws.write_formula(f, 12, f'=A{f + 1}>4', None, f > 4)
        ws.write_number(f, 13, i / 8, porcentaje)
        ws.write_rich_string(f, 14, 'tramo ', negrita, 'negrita', ' y ', rojo, 'rojo')
        if i < 2:
            ws.write_url(f, 15, 'https://example.org/' + str(i), string='Sitio ' + str(i))
    ws.write_string(9, 1, '=no es una fórmula')
    ws.write_string(9, 2, 'x' * 300)
    ws.write_comment(1, 1, 'Un comentario')
    ws.freeze_panes(1, 0)
    ws.autofilter(0, 0, 9, len(encabezado) - 1)
    ws.merge_range('R1:T1', 'Combinadas')
    ws.data_validation('C2:C9', {'validate': 'list', 'source': ['A', 'B']})
    otra = wb.add_worksheet('Segunda')
    otra.write('A1', 'Medida')
    otra.write('B1', 'Valor')
    for i in range(5):
        otra.write_string(i + 1, 0, 'm' + str(i))
        otra.write_number(i + 1, 1, i * i)
    escondida = wb.add_worksheet('Escondida')
    escondida.write('A1', 'nada')
    escondida.hide()
    if not opciones.get('constant_memory'):
        grafico = wb.add_chart({'type': 'column'})
        grafico.add_series({'values': '=Segunda!$B$2:$B$6'})
        hoja_grafico = wb.add_chartsheet('Grafico')
        hoja_grafico.set_chart(grafico)
    wb.close()


# ---------------------------------------------------------------- LibreOffice (UNO)
def libreoffice():
    import uno
    from com.sun.star.beans import PropertyValue

    def prop(nombre, valor):
        p = PropertyValue()
        p.Name, p.Value = nombre, valor
        return p

    perfil = tempfile.mkdtemp(prefix='lo_statsim_')
    tubo = 'statsim_xlsx_%d' % os.getpid()
    proceso = subprocess.Popen(['soffice', '--headless', '--invisible', '--norestore', '--nologo', f'--accept=pipe,name={tubo};urp;',
                                f'-env:UserInstallation=file://{perfil}'], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        local = uno.getComponentContext()
        resolutor = local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver', local)
        ctx = None
        for _ in range(120):
            try:
                ctx = resolutor.resolve(f'uno:pipe,name={tubo};urp;StarOffice.ComponentContext')
                break
            except Exception:
                time.sleep(0.5)
        if ctx is None:
            raise RuntimeError('LibreOffice no respondió')
        escritorio = ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop', ctx)
        doc = escritorio.loadComponentFromURL('private:factory/scalc', '_blank', 0, (prop('Hidden', True),))
        formatos = doc.NumberFormats
        idioma = uno.createUnoStruct('com.sun.star.lang.Locale')

        def formato(codigo):
            k = formatos.queryKey(codigo, idioma, False)
            return k if k != -1 else formatos.addNew(codigo, idioma)

        hojas = doc.Sheets
        h1 = hojas.getByIndex(0)
        h1.Name = 'Hoja1'
        encabezado = ['Participante', 'Edad', 'Ciudad', 'Ingreso', 'Fecha', 'Hora', 'Doble', 'Suma', 'Mayor', 'Division', 'Logico', 'Porcentaje', 'FechaCalc']
        for j, h in enumerate(encabezado):
            h1.getCellByPosition(j, 0).setString(h)
        datos = [('P01', 23, 'Málaga', 1500.5), ('P02', 35, 'Ñuñoa', 2300), ('P03', None, 'São Paulo', 1999.99), ('P04', 41, '', 3100.25), ('P05', 29, 'Zürich', None)]
        f_fecha, f_hora, f_pct = formato('DD/MM/YYYY'), formato('HH:MM:SS'), formato('0.00%')
        for i, (p, edad, ciudad, ingreso) in enumerate(datos, start=1):
            fila = i + 1
            h1.getCellByPosition(0, i).setString(p)
            if edad is not None:
                h1.getCellByPosition(1, i).setValue(edad)
            if ciudad:
                h1.getCellByPosition(2, i).setString(ciudad)
            if ingreso is not None:
                h1.getCellByPosition(3, i).setValue(ingreso)
            c = h1.getCellByPosition(4, i)
            c.setFormula(f'=DATE(2024;{i};{10 + i})')
            c.NumberFormat = f_fecha
            c = h1.getCellByPosition(5, i)
            c.setFormula(f'=TIME({8 + i};{5 * i};{i})')
            c.NumberFormat = f_hora
            h1.getCellByPosition(6, i).setFormula(f'=B{fila}*2')
            h1.getCellByPosition(7, i).setFormula(f'=B{fila}+D{fila}')
            h1.getCellByPosition(8, i).setFormula(f'=IF(B{fila}>30;"mayor";"menor")')
            h1.getCellByPosition(9, i).setFormula(f'=D{fila}/(B{fila}-B{fila})')
            h1.getCellByPosition(10, i).setFormula(f'=B{fila}>30')
            c = h1.getCellByPosition(11, i)
            c.setValue(i / 10)
            c.NumberFormat = f_pct
            c = h1.getCellByPosition(12, i)
            c.setValue(45000 + i)
            c.NumberFormat = f_fecha
        hojas.insertNewByName('Notas', 1)
        h2 = hojas.getByName('Notas')
        h2.getCellByPosition(0, 0).setString('Comentario')
        h2.getCellByPosition(0, 1).setString('Línea uno\nlínea dos')
        h2.getCellByPosition(0, 2).setString('  con espacios  ')
        h2.getCellByPosition(1, 3).setValue(-0.000123)
        hojas.insertNewByName('Oculta', 2)
        h3 = hojas.getByName('Oculta')
        h3.getCellByPosition(0, 0).setValue(99)
        h3.IsVisible = False
        doc.calculateAll()
        url = lambda n: uno.systemPathToFileUrl(ruta(n))
        doc.storeToURL(url('libreoffice.xlsx'), (prop('FilterName', 'Calc MS Excel 2007 XML'),))
        doc.storeToURL(url('libreoffice_ooxml.xlsx'), (prop('FilterName', 'Calc Office Open XML'),))
        doc.storeToURL(url('libreoffice.xls'), (prop('FilterName', 'MS Excel 97'),))
        doc.storeToURL(url('libreoffice.ods'), (prop('FilterName', 'calc8'),))
        doc.storeToURL(url('cifrado.xlsx'), (prop('FilterName', 'Calc MS Excel 2007 XML'), prop('Password', 'secreto')))
        doc.close(True)
        try:
            escritorio.terminate()
        except Exception:
            pass
    finally:
        try:
            proceso.wait(timeout=30)
        except Exception:
            proceso.kill()
        shutil.rmtree(perfil, ignore_errors=True)


# ---------------------------------------------------------------- a mano
NS_MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS_PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'
TIPOS = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>'''

PREFIJOS = {
    '[Content_Types].xml': TIPOS,
    '_rels/.rels': f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="{NS_PKG}"><Relationship Id="rId1" Type="{NS_R}/officeDocument" Target="/xl/workbook.xml"/></Relationships>''',
    'xl/workbook.xml': f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:workbook xmlns:x="{NS_MAIN}" xmlns:r="{NS_R}"><x:workbookPr/><x:bookViews><x:workbookView activeTab="0"/></x:bookViews><x:sheets><x:sheet name="Prefijos &amp; más" sheetId="1" r:id="rId1"/><x:sheet name="Segunda" sheetId="2" state="visible" r:id="rId2"/></x:sheets></x:workbook>''',
    'xl/_rels/workbook.xml.rels': f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="{NS_PKG}"><Relationship Id="rId3" Type="{NS_R}/sharedStrings" Target="sharedStrings.xml"/><Relationship Id="rId1" Type="{NS_R}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="{NS_R}/worksheet" Target="/xl/worksheets/sheet2.xml"/><Relationship Id="rId4" Type="{NS_R}/styles" Target="styles.xml"/></Relationships>''',
    'xl/worksheets/_rels/sheet1.xml.rels': f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="{NS_PKG}"><Relationship Id="rId9" Type="{NS_R}/hyperlink" Target="https://example.org/?a=1&amp;b=2" TargetMode="External"/></Relationships>''',
    'xl/sharedStrings.xml': f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:sst xmlns:x="{NS_MAIN}" count="9" uniqueCount="9"><x:si><x:t>Nombre</x:t></x:si><x:si><x:r><x:rPr><x:b/></x:rPr><x:t>Ri</x:t></x:r><x:r><x:t xml:space="preserve">co </x:t></x:r><x:r><x:t>texto</x:t></x:r></x:si><x:si><x:t>東京</x:t><x:rPh sb="0" eb="2"><x:t>トウキョウ</x:t></x:rPh><x:phoneticPr fontId="1"/></x:si><x:si><x:t>uno_x000D_dos</x:t></x:si><x:si><x:t>literal _x005F_x000D_ sin escapar</x:t></x:si><x:si><x:t>a &amp; b &lt; c &#10;d &#x1F600; &quot;e&quot;</x:t></x:si><x:si/><x:si><x:t/></x:si><x:si><x:t xml:space="preserve">  rodeado  </x:t></x:si></x:sst>''',
    'xl/styles.xml': f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:styleSheet xmlns:x="{NS_MAIN}"><x:numFmts count="2"><x:numFmt numFmtId="164" formatCode="dd/mm/yyyy"/><x:numFmt numFmtId="165" formatCode="&quot;Total: &quot;0.0"/></x:numFmts><x:fonts count="1"><x:font><x:sz val="11"/><x:name val="Calibri"/></x:font></x:fonts><x:fills count="1"><x:fill><x:patternFill patternType="none"/></x:fill></x:fills><x:borders count="1"><x:border/></x:borders><x:cellStyleXfs count="1"><x:xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></x:cellStyleXfs><x:cellXfs count="6"><x:xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><x:xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/><x:xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/><x:xf numFmtId="46" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/><x:xf numFmtId="22" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/><x:xf numFmtId="165" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/></x:cellXfs></x:styleSheet>''',
    'xl/worksheets/sheet1.xml': f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:worksheet xmlns:x="{NS_MAIN}" xmlns:r="{NS_R}"><x:dimension ref="A1:F8"/><x:sheetData>
<x:row r="1"><x:c r="A1" t="s"><x:v>0</x:v></x:c><x:c r="B1" t="inlineStr"><x:is><x:t>Valor</x:t></x:is></x:c><x:c r="C1" t="inlineStr"><x:is><x:r><x:t>Do</x:t></x:r><x:r><x:t>ble</x:t></x:r></x:is></x:c><x:c r="D1" t="str"><x:f>"Fe"&amp;"cha"</x:f><x:v>Fecha</x:v></x:c><x:c r="E1" t="inlineStr"><x:is><x:t>Momento</x:t></x:is></x:c><x:c r="F1" t="inlineStr"><x:is><x:t>Otro</x:t></x:is></x:c></x:row>
<x:row><x:c t="s"><x:v>1</x:v></x:c><x:c><x:v>1.5E-3</x:v></x:c><x:c><x:v>3E-3</x:v></x:c><x:c s="1"><x:v>45000</x:v></x:c><x:c s="4"><x:v>45000.5</x:v></x:c><x:c s="3"><x:v>1.5</x:v></x:c></x:row>
<x:row r="3" spans="1:6"><x:c r="A3" t="s"><x:v>2</x:v></x:c><x:c r="B3"><x:v>-7</x:v></x:c><x:c r="C3"><x:f t="shared" ref="C3:C4" si="0">B3*2</x:f><x:v>-14</x:v></x:c><x:c r="D3" s="2"><x:v>1</x:v></x:c><x:c r="E3" s="4"><x:v>0.25</x:v></x:c><x:c r="F3" s="5"><x:v>12.5</x:v></x:c></x:row>
<x:row r="4"><x:c r="A4" t="s"><x:v>3</x:v></x:c><x:c r="B4" t="b"><x:v>1</x:v></x:c><x:c r="C4"><x:f t="shared" si="0"/><x:v/></x:c><x:c r="D4" s="2"><x:v>59</x:v></x:c><x:c r="E4" t="e"><x:v>#N/A</x:v></x:c><x:c r="F4" t="b"><x:v>0</x:v></x:c></x:row>
<x:row r="6"><x:c r="A6" t="s"><x:v>4</x:v></x:c><x:c r="B6"   s = "0"  ><x:v>  42  </x:v></x:c><x:c r="D6" s="2"><x:v>61</x:v></x:c><x:c r="F6" t="s"><x:v>8</x:v></x:c></x:row>
<x:row r="7"><x:c r="A7" t="s"><x:v>5</x:v></x:c><x:c r="B7" t="s"><x:v>6</x:v></x:c><x:c r="C7" t="s"><x:v>7</x:v></x:c><x:c r="D7" s="1"/><x:c r="E7" t="str"><x:f>A7</x:f><x:v>a &amp; b</x:v></x:c></x:row>
<x:row r="8"><x:c r="A8" t="inlineStr"><x:is><x:t xml:space="preserve">en línea _x0009_ con tabulador</x:t></x:is></x:c><x:c r="B8"><x:v>1E+300</x:v></x:c></x:row>
</x:sheetData><x:hyperlinks><x:hyperlink ref="A1" r:id="rId9"/></x:hyperlinks><x:pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/></x:worksheet>''',
    'xl/worksheets/sheet2.xml': f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="{NS_MAIN}"><sheetData><row r="2"><c r="B2"><v>1</v></c></row><row r="3"><c r="B3" t="inlineStr"><is><t>sin prefijo</t></is></c></row></sheetData></worksheet>''',
}


def escribir_zip(nombre, partes):
    with zipfile.ZipFile(ruta(nombre), 'w', zipfile.ZIP_DEFLATED) as z:
        for parte, contenido in partes.items():
            info = zipfile.ZipInfo(parte, date_time=(2026, 10, 20, 9, 30, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, contenido.encode('utf-8') if isinstance(contenido, str) else contenido)


def a_estricto(partes):
    """La misma libro en la variante «strict» de ISO/IEC 29500: otros espacios de nombres para SpreadsheetML y para los
    tipos de relación (los de OPC y los tipos de contenido no cambian)."""
    salida = {}
    for parte, contenido in partes.items():
        c = contenido.replace(NS_MAIN, 'http://purl.oclc.org/ooxml/spreadsheetml/main')
        c = c.replace(NS_R, 'http://purl.oclc.org/ooxml/officeDocument/relationships')
        if parte == 'xl/workbook.xml':
            c = c.replace('<x:workbook ', '<x:workbook conformance="strict" ')
        salida[parte] = c
    return salida


def a_mayusculas(partes):
    """Las relaciones nombran las partes con otras mayúsculas que el ZIP (OPC no distingue mayúsculas en los nombres)."""
    salida = dict(partes)
    salida['xl/_rels/workbook.xml.rels'] = partes['xl/_rels/workbook.xml.rels'].replace('Target="worksheets/sheet1.xml"', 'Target="Worksheets/Sheet1.XML"').replace('Target="sharedStrings.xml"', 'Target="SharedStrings.xml"')
    return salida


def zip64_con_descriptores(origen, destino):
    """Reescribe un ZIP con lo que hacen algunos programas (Java, Go): cada entrada con descriptor de datos tras ella (los
    tamaños de la cabecera local a cero), los tamaños y la posición del directorio central en el campo extra de ZIP64, y el
    registro de final de 64 bits con su localizador. Python (zipfile) y Excel lo leen."""
    with zipfile.ZipFile(ruta(origen)) as z:
        entradas = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    salida = io.BytesIO()
    central = []
    for nombre, datos in entradas:
        bnombre = nombre.encode('utf-8')
        comp = zlib.compressobj(6, zlib.DEFLATED, -15)
        comprimido = comp.compress(datos) + comp.flush()
        crc = zlib.crc32(datos) & 0xFFFFFFFF
        desplazamiento = salida.tell()
        extra_local = struct.pack('<HHQQ', 1, 16, 0, 0)
        salida.write(struct.pack('<IHHHHHIIIHH', 0x04034b50, 45, 0x0808, 8, 0x4D34, 0x5A94, 0, 0xFFFFFFFF, 0xFFFFFFFF, len(bnombre), len(extra_local)))
        salida.write(bnombre + extra_local + comprimido)
        salida.write(struct.pack('<IIQQ', 0x08074b50, crc, len(comprimido), len(datos)))   # descriptor de ZIP64
        extra = struct.pack('<HHQQQ', 1, 24, len(datos), len(comprimido), desplazamiento)
        central.append(struct.pack('<IHHHHHHIIIHHHHHII', 0x02014b50, 45, 45, 0x0808, 8, 0x4D34, 0x5A94, crc, 0xFFFFFFFF, 0xFFFFFFFF, len(bnombre), len(extra), 0, 0, 0, 0, 0xFFFFFFFF) + bnombre + extra)
    inicio_central = salida.tell()
    for r in central:
        salida.write(r)
    tam_central = salida.tell() - inicio_central
    fin64 = salida.tell()
    salida.write(struct.pack('<IQHHIIQQQQ', 0x06064b50, 44, 45, 45, 0, 0, len(central), len(central), tam_central, inicio_central))
    salida.write(struct.pack('<IIQI', 0x07064b50, 0, fin64, 1))
    salida.write(struct.pack('<IHHHHIIH', 0x06054b50, 0, 0, 0xFFFF, 0xFFFF, 0xFFFFFFFF, 0xFFFFFFFF, 0))
    with open(ruta(destino), 'wb') as f:
        f.write(salida.getvalue())
    with zipfile.ZipFile(ruta(destino)) as z:   # zipfile lo lee entero y comprueba los CRC
        assert z.testzip() is None and [i.filename for i in z.infolist()] == [n for n, _ in entradas]


def otros_formatos():
    escribir_zip('xlsb_falso.xlsb', {'[Content_Types].xml': TIPOS, 'xl/workbook.bin': b'\x83\x01\x00', '_rels/.rels': PREFIJOS['_rels/.rels']})
    escribir_zip('documento_word.docx', {'[Content_Types].xml': TIPOS, 'word/document.xml': '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'})


def main():
    for f in os.listdir(AQUI):
        if re.search(r'\.(xlsx|xls|ods|xlsb|docx)$', f):
            os.remove(ruta(f))
    openpyxl_tipos()
    openpyxl_iso()
    openpyxl_1904()
    xlsxwriter_libro('xlsxwriter_compartidas.xlsx', {})
    xlsxwriter_libro('xlsxwriter_en_linea.xlsx', {'constant_memory': True})
    escribir_zip('a_mano_prefijos.xlsx', PREFIJOS)
    escribir_zip('a_mano_estricto.xlsx', a_estricto(PREFIJOS))
    escribir_zip('a_mano_mayusculas.xlsx', a_mayusculas(PREFIJOS))
    zip64_con_descriptores('openpyxl_tipos.xlsx', 'a_mano_zip64.xlsx')
    otros_formatos()
    libreoffice()
    for f in sorted(os.listdir(AQUI)):
        if not f.endswith('.py'):
            print(f'{f}: {os.path.getsize(ruta(f))} bytes')


if __name__ == '__main__':
    sys.exit(main())
